"""生成 TTS 统一标准技术方案 Word 文档，包含 API_STANDARD 内容。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# ---- 全局样式 ----
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.5
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

for i in range(1, 4):
    h = doc.styles['Heading %d' % i]
    h.font.name = '微软雅黑'
    h.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    h.font.color.rgb = RGBColor(0, 0, 0)


def add_code(doc, text):
    for line in text.strip().split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(line)
        r.font.name = 'Consolas'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_md_table(doc, rows):
    header = [c.strip() for c in rows[0].split('|')[1:-1]]
    data = []
    for r in rows[2:]:
        cells = [c.strip() for c in r.split('|')[1:-1]]
        if cells:
            data.append(cells)
    table = doc.add_table(rows=1 + len(data), cols=len(header))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, t in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = t
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for ri, rd in enumerate(data):
        for ci, t in enumerate(rd):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = t
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


def parse_md(md_path, doc):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    code_buf = []
    tbl_buf = []
    in_code = False
    in_tbl = False

    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith('```'):
            if in_code:
                add_code(doc, '\n'.join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.startswith('|') and line.endswith('|'):
            if not in_tbl:
                in_tbl = True
                tbl_buf = []
            tbl_buf.append(line)
            i += 1
            continue
        else:
            if in_tbl:
                add_md_table(doc, tbl_buf)
                tbl_buf = []
                in_tbl = False

        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.strip() == '':
            pass
        elif line.startswith('- '):
            import re as _re
            m = _re.match(r'- \*\*(.+?)\*\*(.+)', line)
            if m:
                pp = doc.add_paragraph(style='List Bullet')
                rr = pp.add_run(m.group(1))
                rr.bold = True
                pp.add_run(m.group(2))
            else:
                doc.add_paragraph(line[2:], style='List Bullet')
        elif line.strip():
            import re as _re
            p = doc.add_paragraph()
            parts = _re.split(r'(`[^`]+`)', line)
            for part in parts:
                if part.startswith('`') and part.endswith('`'):
                    r2 = p.add_run(part[1:-1])
                    r2.font.name = 'Consolas'
                    r2.font.size = Pt(9)
                else:
                    bolds = _re.split(r'(\*\*[^*]+\*\*)', part)
                    for bp in bolds:
                        if bp.startswith('**') and bp.endswith('**'):
                            r3 = p.add_run(bp[2:-2])
                            r3.bold = True
                        else:
                            p.add_run(bp)
        i += 1

    if in_tbl and tbl_buf:
        add_md_table(doc, tbl_buf)


# ============================================================
# 封面
# ============================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('TTS 统一标准化技术方案')
r.font.size = Pt(22)
r.bold = True
r.font.name = '微软雅黑'
r.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = sub.add_run('文本转语音微服务统一接口规范')
r2.font.size = Pt(14)
r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
r2.font.name = '微软雅黑'
r2.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

doc.add_page_break()

# ============================================================
# 一、项目背景
# ============================================================
doc.add_heading('一、项目背景', level=1)

doc.add_paragraph('当前项目包含三个 TTS（文本转语音）微服务，分别为不同场景提供语音合成能力：')

add_md_table(doc, [
    '| 服务 | 技术栈 | 使用场景 |',
    '|---|---|---|',
    '| TeleAI | FastAPI + 远程 API | 对接 TeleAI 超自然/实时语音 |',
    '| Edge TTS | FastAPI + edge-tts 库 | 微软免费在线语音合成 |',
    '| Sherpa | FastAPI + sherpa-onnx | 本地离线语音合成 |',
])

doc.add_paragraph('三个服务各自独立开发，存在接口不一致、代码重复、认证机制不统一等问题，需要进行统一标准化改造。')

# ============================================================
# 二、问题分析
# ============================================================
doc.add_heading('二、问题分析', level=1)

doc.add_heading('2.1 原状问题', level=2)
doc.add_paragraph('请求参数不统一：不同服务的相同功能使用了不同的字段名和数据结构。')
doc.add_paragraph('错误响应格式不一致：三个服务各自定义了错误码和错误结构。')
doc.add_paragraph('代码大量重复：Pydantic 模型、错误处理、认证逻辑在每个服务中重复定义。')
doc.add_paragraph('认证机制不统一：部分服务存在废弃代码，认证策略不一致。')

doc.add_heading('2.2 统一目标', level=2)
doc.add_paragraph('一次定义，多处复用：抽取共享代码包，消除重复代码。')
doc.add_paragraph('统一请求/响应格式：所有服务使用相同的字段定义和错误结构。')
doc.add_paragraph('统一认证机制：启动认证 + API Key 双层认证，所有端点一致。')
doc.add_paragraph('统一错误码体系：9 种错误码覆盖所有异常场景。')
doc.add_paragraph('向后兼容：旧客户端无需修改即可继续使用。')

# ============================================================
# 三、技术方案
# ============================================================
doc.add_heading('三、技术方案', level=1)

doc.add_heading('3.1 总体架构', level=2)
doc.add_paragraph('核心思想：共享代码下沉到 tts_common/ 包，各服务只保留业务逻辑。所有服务从 tts_common 导入模型、错误处理、认证和工具函数。')

doc.add_heading('3.2 请求模型统一', level=2)
doc.add_paragraph('采用工厂函数模式（create_tts_request_model），各服务调用时传入自身支持的引擎列表、格式列表和默认值，返回统一的 TTSRequest 模型。统一模型包含 11 个字段：text、engine、voice、format、speed、pitch、style、volume、sample_rate、req_id、extra_params。')

doc.add_heading('3.3 认证机制', level=2)
doc.add_paragraph('启动认证：服务启动时要求输入用户名/密码，通过后生成 64 位随机 Token。Docker 等非交互环境自动跳过输入，直接生成 Token。默认账号 admin / tts123。')
doc.add_paragraph('API 认证：所有端点均需在请求头中携带 X-API-Key。接受两种令牌：启动时生成的会话 Token（每次重启更新）和 TTS_API_KEY 环境变量配置的固定 Key（重启不变）。')

doc.add_heading('3.4 错误码体系', level=2)
doc.add_paragraph('统一 9 种错误码覆盖认证、校验、本地合成、上游调用四类场景，所有错误响应使用相同的 JSON 结构。')

doc.add_heading('3.5 代码量对比', level=2)

add_md_table(doc, [
    '| 指标 | 改造前 | 改造后 | 减少 |',
    '|---|---|---|---|',
    '| 单服务重复代码 | ~125 行 | ~0 行 | 100% |',
    '| 三服务总代码量 | ~1500 行 | ~900 行 | ~40% |',
    '| 错误响应格式 | 3 种 | 1 种 | 统一 |',
    '| 认证机制 | 不统一 | 统一 | — |',
])

doc.add_page_break()

# ============================================================
# 四、部署方案
# ============================================================
doc.add_heading('四、部署方案', level=1)

doc.add_heading('4.1 本地开发', level=2)
add_code(doc, 'cd tts/\nuvicorn app:app --host 0.0.0.0 --port 8000')

doc.add_heading('4.2 Docker 部署', level=2)
doc.add_paragraph('使用 docker-compose.yml 一键编排三个服务，支持 docker-compose build && docker-compose up -d。')

add_md_table(doc, [
    '| 服务 | 端口 | 依赖 | 状态 |',
    '|---|---|---|---|',
    '| TeleAI | 8000 | 远程 API | 进行中 |',
    '| Edge TTS | 8001 | 微软服务 | 进行中 |',
    '| Sherpa | 8002 | 本地 ONNX 模型 | 待上传模型 |',
])

doc.add_page_break()

# ============================================================
# 五、API 统一标准（API_STANDARD.md 原文）
# ============================================================
doc.add_heading('五、API 统一标准', level=1)
doc.add_paragraph('以下内容为统一的 HTTP API 标准规范，所有 TTS 服务均遵循此契约。')

parse_md('API_STANDARD.md', doc)

doc.add_page_break()

# ============================================================
# 六、总结
# ============================================================
doc.add_heading('六、总结', level=1)
doc.add_paragraph('本方案通过以下措施实现了三个 TTS 服务的统一标准化：')
doc.add_paragraph('共享代码层：tts_common/ 包消除重复代码，提供统一的模型、错误处理、认证和工具函数。')
doc.add_paragraph('工厂函数模式：一次定义请求模型（11 个统一字段），各服务按需实例化。')
doc.add_paragraph('双层认证：启动认证 + API 认证，所有端点认证策略一致。')
doc.add_paragraph('统一错误码：9 种错误码覆盖全部异常场景，响应格式统一。')
doc.add_paragraph('完整向后兼容：input 别名、extra_params fallback、/text/ 兼容端点。')
doc.add_paragraph('Docker 化部署：docker-compose.yml 一键编排，deploy.ps1 自动部署。')

# ---- 保存 ----
output = 'TTS统一标准技术方案.docx'
doc.save(output)
print('Done:', output)
