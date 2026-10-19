"""Convert API_STANDARD.md to Word document."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import re

doc = Document()

# -- 全局样式 --
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.5
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

for i in range(1, 4):
    h = doc.styles['Heading %d' % i]
    h.font.name = '微软雅黑'
    h.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    h.font.color.rgb = RGBColor(0, 0, 0)


def add_code_block(doc, text):
    for line in text.strip().split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_md_table(doc, rows):
    header_cells = [c.strip() for c in rows[0].split('|')[1:-1]]
    data_rows = []
    for r in rows[2:]:
        cells = [c.strip() for c in r.split('|')[1:-1]]
        if cells:
            data_rows.append(cells)

    table = doc.add_table(rows=1 + len(data_rows), cols=len(header_cells))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, text in enumerate(header_cells):
        cell = table.rows[0].cells[i]
        cell.text = text
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9)

    for r_idx, row_data in enumerate(data_rows):
        for c_idx, text in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


def parse_md_to_docx(md_path, doc):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    code_lines = []
    table_lines = []
    in_code = False
    in_table = False

    while i < len(lines):
        line = lines[i].rstrip()

        # Code block start/end
        if line.startswith('```'):
            if in_code:
                add_code_block(doc, '\n'.join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Table rows
        if line.startswith('|') and line.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                add_md_table(doc, table_lines)
                table_lines = []
                in_table = False

        # Headings
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.strip() == '':
            pass
        elif line.startswith('- **'):
            match = re.match(r'- \*\*(.+?)\*\*(.+)', line)
            if match:
                p = doc.add_paragraph(style='List Bullet')
                r = p.add_run(match.group(1))
                r.bold = True
                p.add_run(match.group(2))
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith('> **'):
            p = doc.add_paragraph(line[4:])
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                run.font.size = Pt(9)
        elif line.strip():
            p = doc.add_paragraph()
            # Inline code: `...`
            parts = re.split(r'(`[^`]+`)', line)
            for part in parts:
                if part.startswith('`') and part.endswith('`'):
                    r = p.add_run(part[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9)
                else:
                    # Bold: **...**
                    bolds = re.split(r'(\*\*[^*]+\*\*)', part)
                    for bp in bolds:
                        if bp.startswith('**') and bp.endswith('**'):
                            r = p.add_run(bp[2:-2])
                            r.bold = True
                        else:
                            p.add_run(bp)
        i += 1

    # trailing table
    if in_table and table_lines:
        add_md_table(doc, table_lines)


parse_md_to_docx('API_STANDARD.md', doc)

output = 'TTS统一API标准.docx'
doc.save(output)
print('Done:', output)
